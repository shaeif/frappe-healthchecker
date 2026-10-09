"""Excel templates for importing clients, AMCs and engineers with Frappe's Data Import tool."""

import io

import frappe
from frappe import _

TEMPLATES = {
	"Client": {
		"sheet": "Clients",
		"fields": [
			("client_code", "C001", "Required, unique (becomes the client ID)."),
			("client_name", "Example Trading WLL", "Required."),
			("account_manager", "am@company.qa", "Email of a user with role AMC Account Manager (default for its AMCs)."),
			("contact_name", "Ahmed Al-Example", ""),
			("contact_designation", "IT Manager", ""),
			("contact_email", "it@example.qa", ""),
			("contact_phone", "+974 5555 0000", ""),
			("cc_emails", "", "Comma separated."),
			("preferred_contact_method", "Email", "Phone, Email, Teams or WhatsApp."),
			("address", "", ""),
			("notes", "", ""),
		],
		"children": [],
	},
	"AMC": {
		"sheet": "AMCs",
		"fields": [
			("client", "C001", "Required: the Client Code of an existing client."),
			("amc_title", "Network AMC", "Optional. Defaults to '<Client> AMC'."),
			("frequency", "Quarterly", "Required: Monthly, Quarterly, Half-yearly or Yearly."),
			("next_due_date", "2026-12-15", "Required. Excel date or YYYY-MM-DD."),
			("contract_start", "2026-01-01", ""),
			("contract_end", "2027-12-31", ""),
			("account_manager", "am@company.qa", "Optional. Defaults to the client's Account Manager."),
			("technical_manager", "", ""),
			("scope", "40 Catalyst switches, 2 C9800 WLCs, 2 PA-3220", ""),
		],
		# child table: (table fieldname, [(child fieldname, example, note)])
		"children": [
			(
				"engineers",
				[
					("engineer", "engineer@company.qa", "One row per engineer: repeat the AMC columns empty on the extra rows."),
					("expertise", "Routing & Switching", "An existing Expertise."),
				],
			)
		],
	},
	"Engineer": {
		"sheet": "Engineers",
		"fields": [
			("user", "engineer@company.qa", "Required: email of an existing user. The profile gives the user the AMC Engineer role."),
			("status", "Active", "Active or Inactive."),
			("mobile_no", "+974 3300 0000", ""),
			("notes", "", ""),
		],
		"children": [
			(
				"expertise",
				[("expertise", "Routing & Switching", "An existing Expertise. One row per extra area: repeat the engineer columns empty on the extra rows.")],
			)
		],
	},
}


@frappe.whitelist()
def download_import_template(doctype: str = "AMC"):
	"""Import template (.xlsx): data sheet (headers only) + a Help sheet."""
	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill

	if doctype not in TEMPLATES:
		frappe.throw(_("No template for {0}").format(doctype))
	if not frappe.has_permission(doctype, "create"):
		frappe.throw(_("Not permitted"), frappe.PermissionError)
	spec = TEMPLATES[doctype]
	meta = frappe.get_meta(doctype)
	headers, help_rows = [], []
	for field, example, note in spec["fields"]:
		headers.append(meta.get_label(field))
		help_rows.append((meta.get_label(field), example, _(note) if note else ""))
	for table_field, child_fields in spec["children"]:
		table_label = meta.get_label(table_field)
		child_meta = frappe.get_meta(meta.get_field(table_field).options)
		for field, example, note in child_fields:
			header = f"{child_meta.get_label(field)} ({table_label})"
			headers.append(header)
			help_rows.append((header, example, _(note) if note else ""))

	wb = Workbook()
	ws = wb.active
	ws.title = spec["sheet"]
	ws.append(headers)
	for cell in ws[1]:
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = PatternFill("solid", fgColor="0B4F8A")
	for i, header in enumerate(headers, start=1):
		ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, len(header) + 4)

	help_ws = wb.create_sheet("Help")
	help_ws.append([_("Column"), _("Example"), _("Notes")])
	for row in help_rows:
		help_ws.append(list(row))
	help_ws.append([])
	help_ws.append([_("Fill the '{0}' sheet only, then use {1} list > Import > Import from Excel.").format(spec["sheet"], _(doctype))])
	for col, width in (("A", 30), ("B", 44), ("C", 70)):
		help_ws.column_dimensions[col].width = width

	buffer = io.BytesIO()
	wb.save(buffer)
	frappe.response["filename"] = f"{frappe.scrub(doctype)}_import_template.xlsx"
	frappe.response["filecontent"] = buffer.getvalue()
	frappe.response["type"] = "binary"
