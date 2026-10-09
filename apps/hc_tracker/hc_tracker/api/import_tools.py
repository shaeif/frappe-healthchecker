"""Excel template for importing contracts with Frappe's Data Import tool."""

import io

import frappe
from frappe import _

# (fieldname, example value) - column headers are the field labels, which Data Import maps automatically
TEMPLATE_FIELDS = [
	("client_id", "C001"),
	("client_name", "Example Trading WLL"),
	("frequency", "Quarterly"),
	("next_due_date", "2026-12-15"),
	("assigned_engineer", "engineer@company.qa"),
	("account_manager", "am@company.qa"),
	("technical_manager", ""),
	("helpdesk_contact", ""),
	("contract_start", "2026-01-01"),
	("contract_end", "2027-12-31"),
	("scope", "40 Catalyst switches, 2 C9800 WLCs, 2 PA-3220"),
	("client_contact_name", "Ahmed Al-Example"),
	("client_contact_email", "it@example.qa"),
	("client_contact_phone", "+974 5555 0000"),
	("client_cc_emails", ""),
	("preferred_contact_method", "Email"),
	("notes", ""),
]


@frappe.whitelist()
def download_import_template():
	"""HC Contract import template (.xlsx): sheet 'Contracts' (headers only, imported) and sheet 'Help'."""
	from openpyxl import Workbook
	from openpyxl.styles import Font, PatternFill

	if not frappe.has_permission("HC Contract", "create"):
		frappe.throw(_("Not permitted"), frappe.PermissionError)
	meta = frappe.get_meta("HC Contract")
	wb = Workbook()
	ws = wb.active
	ws.title = "Contracts"
	headers = [meta.get_label(f) for f, _example in TEMPLATE_FIELDS]
	ws.append(headers)
	for cell in ws[1]:
		cell.font = Font(bold=True, color="FFFFFF")
		cell.fill = PatternFill("solid", fgColor="0B5CAD")
	for i, header in enumerate(headers, start=1):
		ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = max(14, len(header) + 4)

	help_ws = wb.create_sheet("Help")
	help_ws.append([_("Column"), _("Example"), _("Notes")])
	notes = {
		"client_id": _("Required, unique (becomes the contract ID)."),
		"client_name": _("Required."),
		"frequency": _("Required: Monthly, Quarterly, Half-yearly or Yearly."),
		"next_due_date": _("Required. Excel date or YYYY-MM-DD."),
		"assigned_engineer": _("Email of an existing user with role HC Engineer."),
		"account_manager": _("Email of an existing user with role HC Account Manager."),
		"preferred_contact_method": _("Phone, Email, Teams or WhatsApp."),
	}
	for field, example in TEMPLATE_FIELDS:
		help_ws.append([meta.get_label(field), example, notes.get(field, "")])
	help_ws.append([])
	help_ws.append([_("Fill the 'Contracts' sheet only (one row per client), then use HC Contract list > Import from Excel.")])
	for col, width in (("A", 26), ("B", 44), ("C", 60)):
		help_ws.column_dimensions[col].width = width

	buffer = io.BytesIO()
	wb.save(buffer)
	frappe.response["filename"] = "hc_contract_import_template.xlsx"
	frappe.response["filecontent"] = buffer.getvalue()
	frappe.response["type"] = "binary"
